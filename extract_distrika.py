import openpyxl
import json

# Ouvrir le fichier Excel
wb = openpyxl.load_workbook('FKT.xlsx', read_only=True)
ws = wb.active

# Trouver la colonne DISTRIKA
distrika_col = None
for col in range(1, ws.max_column + 1):
    if ws.cell(1, col).value == 'DISTRIKA':
        distrika_col = col
        break

if distrika_col:
    # Extraire les valeurs uniques
    distrikas = set()
    for row in range(2, ws.max_row + 1):
        value = ws.cell(row, distrika_col).value
        if value and str(value).strip():
            distrikas.add(str(value).strip())
    
    # Trier et afficher
    sorted_distrikas = sorted(list(distrikas))
    print(json.dumps(sorted_distrikas, ensure_ascii=False, indent=2))
else:
    print("Colonne DISTRIKA non trouvée")

wb.close()
